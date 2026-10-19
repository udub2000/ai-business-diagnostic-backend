"""Extracted legacy financial parser functions from app/main.py for isolated staging.
Do not import app.main (database/email side effects). Review changes against original.
"""
from __future__ import annotations
import csv
import re
from pathlib import Path
from typing import Any, Optional

def parse_money_or_number(value: Any) -> Optional[float]:
    """Parse accounting-style numbers safely."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    is_negative = text.startswith("(") and text.endswith(")")
    cleaned = re.sub(r"[^0-9.\-]", "", text)
    if cleaned in {"", "-", ".", "-."} or cleaned.count(".") > 1:
        return None
    try:
        number = float(cleaned)
        return -abs(number) if is_negative else number
    except ValueError:
        return None

def normalize_label(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()

def rows_from_csv_or_text(path: Path) -> list[list[Any]]:
    rows: list[list[Any]] = []
    text = path.read_text(encoding="utf-8", errors="ignore")
    if "<table" in text.lower() or "<tr" in text.lower():
        tr_blocks = re.findall(r"<tr[^>]*>(.*?)</tr>", text, flags=re.IGNORECASE | re.DOTALL)
        for tr in tr_blocks:
            cells = re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", tr, flags=re.IGNORECASE | re.DOTALL)
            cleaned_cells = [re.sub(r"<[^>]+>", "", c).strip() for c in cells]
            if any(cleaned_cells):
                rows.append(cleaned_cells)
        return rows
    sample = text[:2048]
    delimiter = "\t" if "\t" in sample and sample.count("\t") > sample.count(",") else ","
    reader = csv.reader(text.splitlines(), delimiter=delimiter)
    for row in reader:
        if any(str(cell).strip() for cell in row):
            rows.append(row)
    return rows

def rows_from_xlsx(path: Path) -> list[list[Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []
    wb = load_workbook(path, data_only=True, read_only=True)
    rows: list[list[Any]] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            values = list(row)
            if any(v not in (None, "") for v in values):
                rows.append(values)
    return rows

def extract_pdf_text(path: Path) -> str:
    text = ""
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        for page in reader.pages:
            text += "\n" + (page.extract_text() or "")
        if text.strip():
            return text
    except Exception:
        pass
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(str(path))
        for page in reader.pages:
            text += "\n" + (page.extract_text() or "")
    except Exception:
        return ""
    return text

def split_pdf_label_and_values(line: str) -> Optional[list[Any]]:
    cleaned = re.sub(r"\s+", " ", line.strip())
    if not cleaned:
        return None
    money_pattern = r"\(?\$?\s*-?\d[\d,]*(?:\.\d+)?\)?"
    matches = list(re.finditer(money_pattern, cleaned))
    if not matches:
        return [cleaned]
    first_number = matches[0]
    label = cleaned[: first_number.start()].strip(" -:\t")
    if not label:
        return None
    values = [m.group(0).strip() for m in matches]
    trailing = cleaned[matches[-1].end():].strip(" -:\t")
    row: list[Any] = [label, *values]
    if trailing:
        row.append(trailing)
    return row

def rows_from_pdf_text_lines(text: str) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for raw_line in text.splitlines():
        parsed = split_pdf_label_and_values(raw_line.strip())
        if parsed and len(parsed) > 1:
            rows.append(parsed)
    return rows

def rows_from_pdf_column_text(text: str) -> list[list[Any]]:
    lines = [re.sub(r"\s+", " ", line.strip()) for line in text.splitlines() if line.strip()]
    rows: list[list[Any]] = []
    money_pattern = re.compile(r"^\(?\$?\s*-?\d[\d,]*(?:\.\d+)?\)?$")
    i = 0
    while i < len(lines):
        if not metric_key_from_label(lines[i]):
            i += 1
            continue
        values: list[str] = []
        j = i + 1
        while j < len(lines) and len(values) < 5:
            if money_pattern.match(lines[j]):
                values.append(lines[j]); j += 1; continue
            if metric_key_from_label(lines[j]):
                break
            if values:
                break
            j += 1
        if values:
            rows.append([lines[i], *values]); i = j
        else:
            i += 1
    return rows

def rows_from_pdf(path: Path) -> list[list[Any]]:
    text = extract_pdf_text(path)
    if not text.strip():
        return []
    rows = rows_from_pdf_text_lines(text)
    recognized_rows = [row for row in rows if row and metric_key_from_label(str(row[0]))]
    if len(recognized_rows) < 3:
        column_rows = rows_from_pdf_column_text(text)
        if len(column_rows) > len(recognized_rows):
            rows = column_rows
    return rows

def rows_from_financial_file(path: Path) -> list[list[Any]]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".xls"}:
        return rows_from_csv_or_text(path)
    if suffix == ".xlsx":
        return rows_from_xlsx(path)
    if suffix == ".pdf":
        return rows_from_pdf(path)
    return []

def looks_like_year_number(number: float) -> bool:
    return 1900 <= abs(number) <= 2100 and float(number).is_integer()

def cell_year(value: Any) -> Optional[int]:
    text = str(value or "").strip()
    match = re.search(r"\b(19\d{2}|20\d{2}|21\d{2})\b", text)
    return int(match.group(1)) if match else None

def cell_has_financial_number(value: Any) -> bool:
    number = parse_money_or_number(value)
    return number is not None and not looks_like_year_number(number)

FINANCIAL_LABEL_ALIASES: dict[str, list[str]] = {
    "revenue": ["total revenue", "service revenue", "revenue", "sales", "net sales"],
    "cogs": ["cost of goods sold", "cogs", "cost of sales", "cost of revenue", "total cost of revenue", "direct costs"],
    "gross_profit": ["gross profit"],
    "operating_expenses": ["total operating expenses", "operating expenses", "opex", "total expenses"],
    "operating_income": ["operating income", "income from operations", "ebit"],
    "interest_expense": ["interest expense"],
    "net_income": ["net income", "net profit"],
    "cash": ["cash", "cash and cash equivalents", "cash equivalents", "ending cash"],
    "accounts_receivable": ["accounts receivable", "trade receivables", "ar", "a r"],
    "inventory": ["inventory"],
    "total_assets": ["total assets"],
    "accounts_payable": ["accounts payable", "trade payables", "ap", "a p"],
    "debt": ["total debt", "current debt", "long term debt", "long-term debt", "debt", "loans", "notes payable"],
    "owner_equity": ["owner equity", "owners equity", "total equity", "equity", "retained earnings"],
    "net_cash_from_operations": ["net cash from operations", "net cash provided by operating activities", "operating cash flow", "cash from operations"],
    "capital_expenditures": ["capital expenditures", "capex", "purchases of property and equipment"],
    "debt_repayment": ["debt repayment", "loan repayment", "principal repayment", "repayments of debt"],
    "net_change_in_cash": ["net change in cash", "change in cash", "net increase in cash", "net decrease in cash"],
}

STATEMENT_METRIC_RULES: dict[str, dict[str, dict[str, Any]]] = {
    "income_statement": {
        "revenue": {"aliases": ["total revenue", "net sales", "total sales", "revenue", "service revenue"], "exclude": ["materials revenue", "product revenue"], "prefer_total": True},
        "cogs": {"aliases": ["cost of goods sold", "total cost of revenue", "cost of revenue", "cost of sales", "cogs", "direct costs"], "exclude": ["direct labor", "materials and supplies", "subcontractors"], "prefer_total": True},
        "gross_profit": {"aliases": ["gross profit"], "exclude": [], "prefer_total": False},
        "operating_expenses": {"aliases": ["total operating expenses", "operating expenses", "opex"], "exclude": ["office payroll", "rent", "marketing", "insurance", "software"], "prefer_total": True},
        "operating_income": {"aliases": ["operating income", "income from operations", "ebit"], "exclude": [], "prefer_total": False},
        "interest_expense": {"aliases": ["interest expense"], "exclude": [], "prefer_total": False},
        "net_income": {"aliases": ["net income", "net profit"], "exclude": [], "prefer_total": False},
    },
    "balance_sheet": {
        "cash": {"aliases": ["cash", "cash and cash equivalents", "cash equivalents"], "exclude": [], "prefer_total": False},
        "accounts_receivable": {"aliases": ["accounts receivable", "trade receivables", "a r", "ar"], "exclude": ["change in accounts receivable"], "prefer_total": False},
        "inventory": {"aliases": ["inventory"], "exclude": ["change in inventory"], "prefer_total": False},
        "total_assets": {"aliases": ["total assets"], "exclude": ["total current assets"], "prefer_total": True},
        "accounts_payable": {"aliases": ["accounts payable", "trade payables", "a p", "ap"], "exclude": ["change in accounts payable"], "prefer_total": False},
        "debt": {"aliases": ["total debt", "current debt", "long term debt", "long-term debt", "notes payable", "loans"], "exclude": ["debt repayment", "debt issued"], "prefer_total": True, "aggregate_components": True},
        "owner_equity": {"aliases": ["owner equity", "owners equity", "total equity", "equity", "retained earnings"], "exclude": [], "prefer_total": False},
    },
    "cash_flow_statement": {
        "net_income": {"aliases": ["net income"], "exclude": [], "prefer_total": False},
        "net_cash_from_operations": {"aliases": ["net cash from operations", "net cash provided by operating activities", "operating cash flow", "cash from operations"], "exclude": [], "prefer_total": True},
        "capital_expenditures": {"aliases": ["capital expenditures", "capex", "purchases of property and equipment"], "exclude": [], "prefer_total": False},
        "debt_repayment": {"aliases": ["debt repayment", "loan repayment", "principal repayment", "repayments of debt"], "exclude": ["debt issued"], "prefer_total": False},
        "net_change_in_cash": {"aliases": ["net change in cash", "change in cash", "net increase in cash", "net decrease in cash"], "exclude": [], "prefer_total": False},
        "cash": {"aliases": ["ending cash", "cash at end of period"], "exclude": ["net change"], "prefer_total": False},
    },
}

def metric_key_from_label(label: str) -> Optional[str]:
    normalized = normalize_label(label)
    if not normalized:
        return None
    best_key: Optional[str] = None
    best_score = -1
    for key, aliases in FINANCIAL_LABEL_ALIASES.items():
        for alias in aliases:
            normalized_alias = normalize_label(alias)
            score = -1
            if normalized == normalized_alias:
                score = 1000 + len(normalized_alias)
            elif normalized.endswith(normalized_alias):
                score = 700 + len(normalized_alias)
            elif normalized_alias in normalized:
                score = 400 + len(normalized_alias)
            if score > best_score:
                best_key, best_score = key, score
    return best_key

def detect_table_structure(rows: list[list[Any]]) -> dict[str, Any]:
    best_header_index: Optional[int] = None
    best_year_columns: dict[int, int] = {}
    best_score = -1
    for idx, row in enumerate(rows[:15]):
        year_columns = {col_idx: cell_year(cell) for col_idx, cell in enumerate(row) if cell_year(cell) is not None}
        text_cells = sum(1 for cell in row if normalize_label(cell) and not cell_year(cell))
        score = len(year_columns) * 10 + text_cells
        if year_columns and score > best_score:
            best_header_index, best_year_columns, best_score = idx, year_columns, score
    if best_year_columns:
        latest_year_col = max(best_year_columns, key=lambda c: best_year_columns[c])
        first_value_col = min(best_year_columns)
        return {"header_row": best_header_index, "year_columns": best_year_columns, "latest_value_column": latest_year_col, "latest_year": best_year_columns[latest_year_col], "label_columns": list(range(first_value_col)) or [0], "single_period": False}
    label_width_counts: dict[int, int] = {}
    for row in rows:
        numeric_cols = [i for i, cell in enumerate(row) if cell_has_financial_number(cell)]
        if numeric_cols and min(numeric_cols) > 0:
            label_width_counts[min(numeric_cols)] = label_width_counts.get(min(numeric_cols), 0) + 1
    label_width = max(label_width_counts, key=label_width_counts.get) if label_width_counts else 1
    return {"header_row": None, "year_columns": {}, "latest_value_column": None, "latest_year": None, "label_columns": list(range(max(1, label_width))), "single_period": True}

def row_label_parts(row: list[Any], structure: dict[str, Any]) -> tuple[str, str, str]:
    labels = [str(row[i]).strip() for i in (structure.get("label_columns") or [0]) if i < len(row) and str(row[i] or "").strip()]
    category = labels[0] if len(labels) > 1 else ""
    line_item = labels[-1] if labels else ""
    return category, line_item, " ".join(labels).strip()

def value_for_row(row: list[Any], structure: dict[str, Any]) -> tuple[Optional[float], Optional[int], Optional[str], float]:
    latest_col = structure.get("latest_value_column")
    latest_year = structure.get("latest_year")
    if latest_col is not None and latest_col < len(row):
        value = parse_money_or_number(row[latest_col])
        if value is not None and not looks_like_year_number(value):
            return value, latest_col, str(latest_year) if latest_year else None, 0.98
    numeric_cells: list[tuple[int, float]] = []
    max_label_col = max(structure.get("label_columns") or [0])
    for idx, cell in enumerate(row):
        if idx <= max_label_col:
            continue
        number = parse_money_or_number(cell)
        if number is not None and not looks_like_year_number(number):
            numeric_cells.append((idx, number))
    if numeric_cells:
        col_idx, value = numeric_cells[-1]
        return value, col_idx, None, 0.72 if structure.get("single_period") else 0.62
    return None, None, None, 0.0

def alias_match_score(label: str, aliases: list[str], exclusions: list[str]) -> tuple[float, Optional[str]]:
    normalized = normalize_label(label)
    if not normalized:
        return 0.0, None
    for exclusion in exclusions:
        normalized_exclusion = normalize_label(exclusion)
        if normalized_exclusion and normalized_exclusion in normalized:
            return 0.0, None
    best_score, best_alias = 0.0, None
    for alias in aliases:
        normalized_alias = normalize_label(alias)
        if normalized == normalized_alias:
            score = 1.0
        elif normalized.endswith(normalized_alias):
            score = 0.92
        elif normalized_alias in normalized:
            score = 0.80
        else:
            score = 0.0
        if score > best_score:
            best_score, best_alias = score, alias
    return best_score, best_alias

def confidence_for_candidate(alias_score: float, value_confidence: float, statement_type: str, metric_key: str, category: str, line_item: str, structure: dict[str, Any], rule: dict[str, Any]) -> float:
    confidence = alias_score * 0.65 + value_confidence * 0.25
    if structure.get("latest_year") is not None:
        confidence += 0.07
    if statement_type in STATEMENT_METRIC_RULES:
        confidence += 0.05
    normalized_line = normalize_label(line_item)
    normalized_category = normalize_label(category)
    if rule.get("prefer_total") and "total" in normalized_line:
        confidence += 0.06
    if metric_key == "debt" and any(x in normalized_line for x in ["debt", "loan", "notes payable"]):
        confidence += 0.04
    if metric_key == "total_assets" and normalized_line == "total assets":
        confidence += 0.08
    if metric_key == "revenue" and normalized_category == "revenue" and "total" in normalized_line:
        confidence += 0.06
    return round(min(confidence, 0.99), 3)

def extract_statement_metrics_from_rows(rows: list[list[Any]], statement_type: str, file_name: str = "") -> dict[str, Any]:
    structure = detect_table_structure(rows)
    rules = STATEMENT_METRIC_RULES.get(statement_type, {})
    candidates_by_metric: dict[str, list[dict[str, Any]]] = {key: [] for key in rules.keys()}
    debug: dict[str, Any] = {}
    row_start = (structure.get("header_row") + 1) if structure.get("header_row") is not None else 0
    for row_index, row in enumerate(rows[row_start:], start=row_start + 1):
        if not row:
            continue
        category, line_item, combined_label = row_label_parts(row, structure)
        label_for_matching = line_item or combined_label
        if not normalize_label(label_for_matching):
            continue
        value, value_col, column_used, value_confidence = value_for_row(row, structure)
        if value is None:
            continue
        for metric_key, rule in rules.items():
            alias_score, alias_used = alias_match_score(label_for_matching, rule.get("aliases", []), rule.get("exclude", []))
            if alias_score <= 0:
                alias_score, alias_used = alias_match_score(combined_label, rule.get("aliases", []), rule.get("exclude", []))
            if alias_score <= 0:
                continue
            confidence = confidence_for_candidate(alias_score, value_confidence, statement_type, metric_key, category, line_item, structure, rule)
            candidates_by_metric.setdefault(metric_key, []).append({
                "metric": metric_key, "value": value, "source_file": file_name, "row": row_index,
                "label_used": label_for_matching, "category": category, "alias_used": alias_used,
                "value_selected": value, "value_column_index": value_col,
                "column_used": column_used or ("single_period_or_rightmost_value" if structure.get("single_period") else None),
                "confidence": confidence,
            })
    metrics: dict[str, float] = {}
    unavailable: dict[str, Any] = {}
    threshold = 0.78
    for metric_key, candidates in candidates_by_metric.items():
        if not candidates:
            unavailable[metric_key] = {"reason": "not_found", "confidence": 0.0, "source_file": file_name}
            continue
        candidates.sort(key=lambda c: (c["confidence"], "total" in normalize_label(c.get("label_used")), abs(float(c.get("value") or 0)) if rules.get(metric_key, {}).get("prefer_total") else 0), reverse=True)
        best = candidates[0]
        if statement_type == "balance_sheet" and metric_key == "debt":
            total_debt_candidates = [c for c in candidates if "total debt" in normalize_label(c.get("label_used"))]
            component_candidates = [c for c in candidates if any(x in normalize_label(c.get("label_used")) for x in ["current debt", "long term debt", "long term debt", "long term debt"]) and c["confidence"] >= threshold]
            if total_debt_candidates and total_debt_candidates[0]["confidence"] >= threshold:
                best = total_debt_candidates[0]
            elif len(component_candidates) >= 2:
                component_sum = sum(float(c["value"]) for c in component_candidates)
                best = {"metric": "debt", "value": component_sum, "source_file": file_name, "row": [c["row"] for c in component_candidates], "label_used": " + ".join(c["label_used"] for c in component_candidates), "category": "Current Liabilities / Long-Term Liabilities", "alias_used": "debt components", "value_selected": component_sum, "value_column_index": [c["value_column_index"] for c in component_candidates], "column_used": component_candidates[0].get("column_used"), "confidence": round(min(min(c["confidence"] for c in component_candidates) + 0.03, 0.96), 3), "aggregation": "sum_of_current_debt_and_long_term_debt"}
        if best["confidence"] >= threshold:
            metrics[metric_key] = float(best["value"])
            debug[metric_key] = best
        else:
            unavailable[metric_key] = {"reason": "low_confidence", "source_file": file_name, "row": best.get("row"), "label_used": best.get("label_used"), "value_selected": best.get("value_selected"), "confidence": best.get("confidence", 0.0)}
    return {"metrics": metrics, "debug": debug, "unavailable_metrics": unavailable, "structure": structure}

def extract_metrics_from_rows(rows: list[list[Any]], statement_type: str = "unknown", file_name: str = "") -> dict[str, float]:
    return extract_statement_metrics_from_rows(rows, statement_type, file_name).get("metrics", {})

def derive_financial_metrics(metrics: dict[str, Any]) -> dict[str, Any]:
    derived = dict(metrics)
    revenue = parse_money_or_number(derived.get("revenue"))
    cogs = parse_money_or_number(derived.get("cogs"))
    gross_profit = parse_money_or_number(derived.get("gross_profit"))
    net_income = parse_money_or_number(derived.get("net_income"))
    cash = parse_money_or_number(derived.get("cash"))
    debt = parse_money_or_number(derived.get("debt"))
    net_cash_from_operations = parse_money_or_number(derived.get("net_cash_from_operations"))
    if gross_profit is None and revenue is not None and cogs is not None:
        gross_profit = revenue - cogs
        derived["gross_profit"] = gross_profit
    if revenue and gross_profit is not None:
        derived["gross_margin_pct"] = round((gross_profit / revenue) * 100, 1)
    if revenue:
        derived["monthly_revenue"] = round(revenue / 12, 2)
    if revenue and net_income is not None:
        derived["net_margin_pct"] = round((net_income / revenue) * 100, 1)
    if cash is not None and debt is not None:
        derived["cash_to_debt_ratio"] = round(cash / debt, 2) if debt else None
    if net_cash_from_operations is not None and net_cash_from_operations < 0:
        derived["cash_stress_signal"] = "high"
    elif cash is not None and debt is not None and debt > cash * 2:
        derived["cash_stress_signal"] = "moderate"
    elif cash is not None:
        derived["cash_stress_signal"] = "low"
    return derived

def parse_financial_uploads(uploads: dict[str, Optional[dict[str, str]]]) -> dict[str, Any]:
    parsed: dict[str, Any] = {"statements": {}, "combined_metrics": {}, "financial_parse_debug": {}, "unavailable_metrics": {}, "parse_errors": []}
    for statement_type, metadata in uploads.items():
        if not metadata:
            continue
        path = Path(metadata["file_path"])
        try:
            rows = rows_from_financial_file(path)
            extraction = extract_statement_metrics_from_rows(rows, statement_type, metadata.get("file_name", path.name))
            metrics = extraction.get("metrics", {})
            debug = extraction.get("debug", {})
            unavailable = extraction.get("unavailable_metrics", {})
            parsed["statements"][statement_type] = {"file_name": metadata["file_name"], "row_count": len(rows), "metrics": metrics, "financial_parse_debug": debug, "unavailable_metrics": unavailable, "table_structure": extraction.get("structure", {})}
            parsed["combined_metrics"].update(metrics)
            parsed["financial_parse_debug"].update({f"{statement_type}.{k}": v for k, v in debug.items()})
            parsed["unavailable_metrics"].update({f"{statement_type}.{k}": v for k, v in unavailable.items()})
        except Exception as exc:
            parsed["parse_errors"].append(f"{statement_type}: {exc}")
    parsed["combined_metrics"] = derive_financial_metrics(parsed["combined_metrics"])
    return parsed
